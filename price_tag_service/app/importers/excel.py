from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import sessionmaker

from app.models.entities import ImportJob
from app.schemas.tasks import PhotoCreate
from app.services.task_service import create_task, refresh_task_counters

logger = logging.getLogger(__name__)

FIELD_MAP = {
    "id": "source_row_id",
    "storeCode": "store_code",
    "storeName": "store_name",
    "storeRegion": "store_region",
    "firstChannel": "channel_level1",
    "secondChannel": "channel_level2",
    "xlOrderId": "xl_order_id",
    "xlMarketId": "xl_market_id",
    "visitStartTime": "visit_start_time",
    "visitFinishTime": "visit_finish_time",
}


def _cell_value(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def extract_photos(
    excel_path: Path,
    *,
    photo_sources: list[str],
    max_photos: int,
) -> tuple[list[PhotoCreate], int, int]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        if worksheet is None:
            return [], 0, 0
        worksheet.reset_dimensions()
        headers = {
            _cell_value(cell.value): column_index
            for column_index, cell in enumerate(next(worksheet.iter_rows(min_row=1, max_row=1)))
        }
        required_columns = list(FIELD_MAP) + ["deletedAt"] + photo_sources
        missing_columns = [column for column in required_columns if column not in headers]
        if missing_columns:
            raise ValueError(f"Excel 缺少必要字段: {', '.join(missing_columns)}")

        photos: list[PhotoCreate] = []
        total_rows = 0
        valid_rows = 0
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            total_rows += 1
            if _cell_value(row[headers["deletedAt"]]):
                continue
            source_metadata = {
                target: _cell_value(row[headers[source]])
                for source, target in FIELD_MAP.items()
                if _cell_value(row[headers[source]])
            }
            row_has_photo = False
            for photo_source in photo_sources:
                raw_urls = _cell_value(row[headers[photo_source]])
                if not raw_urls:
                    continue
                for image_url in raw_urls.split(","):
                    image_url = image_url.strip()
                    if not image_url:
                        continue
                    row_has_photo = True
                    if len(photos) < max_photos:
                        photos.append(
                            PhotoCreate(
                                image_url=image_url,
                                source_row_id=source_metadata.get("source_row_id"),
                                source_photo_type=photo_source,
                                image_name=image_url.rsplit("/", 1)[-1][:512],
                                source_metadata=source_metadata,
                            )
                        )
            if row_has_photo:
                valid_rows += 1
        return photos, total_rows, valid_rows
    finally:
        workbook.close()


def run_import_job(
    *,
    session_factory: sessionmaker,
    import_job_id: str,
    excel_path: Path,
    photo_sources: list[str],
    max_photos: int,
) -> None:
    with session_factory() as session:
        job = session.get(ImportJob, import_job_id)
        if job is None:
            return
        job.status = "RUNNING"
        session.commit()
        try:
            photos, total_rows, valid_rows = extract_photos(
                excel_path,
                photo_sources=photo_sources,
                max_photos=max_photos,
            )
            task, created = create_task(
                session,
                source_type="EXCEL_IMPORT",
                photos=photos,
                task_config={"photo_sources": photo_sources, "max_photos": max_photos},
                created_by="excel-import",
            )
            job.task_id = task.id
            job.total_rows = total_rows
            job.valid_rows = valid_rows
            job.created_photos = created
            job.skipped_photos = max(0, len(photos) - created)
            job.status = "COMPLETED"
            refresh_task_counters(session, task.id)
            session.commit()
        except Exception as exc:
            session.rollback()
            job = session.get(ImportJob, import_job_id)
            if job is not None:
                job.status = "FAILED"
                job.error_message = str(exc)[:4000]
                session.commit()
            logger.exception("Excel import job %s failed", import_job_id)
